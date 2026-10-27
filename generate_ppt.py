from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation(output_path):
    prs = Presentation()

    # Define Colors
    color_primary = RGBColor(249, 166, 2)  # Orange
    color_bg = RGBColor(13, 13, 13)       # Dark

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # --- SLIDE 1: COVER ---
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, color_bg)

    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "AGROSYNC AI"
    subtitle.text = "The Digital Backbone of Agri-Export\nDigitizing Post-Harvest Value Chains with Edge AI & IoT"

    # Style Title
    title.text_frame.paragraphs[0].font.color.rgb = color_primary
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True

    # --- SLIDE 2: THE PROBLEM ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, color_bg)

    title = slide.shapes.title
    title.text = "The $13 Billion Blind Spot"
    title.text_frame.paragraphs[0].font.color.rgb = color_primary

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "The Post-Harvest Crisis in India:"
    
    p = tf.add_paragraph()
    p.text = "    • 35-40% Horticulture Loss – Due to blind storage & inefficient grading."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "    • High Rejection Rates – 15%+ in global markets due to subjective manual grading."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "The Opportunity: 1.5x Farmer Income by digitizing the value chain."
    p.level = 0
    p.font.color.rgb = color_primary

    # --- SLIDE 3: HARDWARE CORE ---
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, color_bg)
    title = slide.shapes.title
    title.text = "Industrial Hardware Architecture"
    title.text_frame.paragraphs[0].font.color.rgb = color_primary

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Edge-Intelligence (CPS) Integration:"
    
    p = tf.add_paragraph()
    p.text = "    • Vision Node: ESP32-S3 + OV2640 Lens (On-device TFLite Micro grading)."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "    • Gas Sensory Hub: SGP30 (Ethylene proxy) + MQ-3 (Fermentation alcohols)."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "    • Power & Comms: Solar + Li-ion power | LoRaWAN / MQTT Connectivity."
    p.level = 1

    # --- SLIDE 4: TECH STACK ---
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, color_bg)
    title = slide.shapes.title
    title.text = "Edge-Heavy Tech Stack"
    title.text_frame.paragraphs[0].font.color.rgb = color_primary

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Built for Rural Scalability:"
    
    p = tf.add_paragraph()
    p.text = "    • AI Core: INT8 Quantized MobileNetV2 (Real-time Grading)."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "    • Backend: FastAPI (High Perf) + InfluxDB (Time-series Sensor Logs)."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "    • Trust Layer: Hyperledger Blockchain for immutable export traceability."
    p.level = 1

    # --- SLIDE 5: VALUE CHAIN ---
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, color_bg)
    title = slide.shapes.title
    title.text = "The Value Chain Impact"
    title.text_frame.paragraphs[0].font.color.rgb = color_primary

    content = slide.placeholders[1]
    tf = content.text_frame
    
    p = tf.add_paragraph()
    p.text = "Farmer → FPO → Exporter → Global Hubs"
    p.font.bold = True
    p.font.size = Pt(28)

    p = tf.add_paragraph()
    p.text = "    • Standardized Grading: Fair pricing for farmers within seconds."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "    • Spoilage Alerts: Predicting decay 72 hours before it occurs."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "    • Global Trust: Blockchain certified batches for premium markets."
    p.level = 1

    # --- SLIDE 6: WHY AGROSYNC? ---
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, color_bg)
    title = slide.shapes.title
    title.text = "Why AgroSync Wins?"
    title.text_frame.paragraphs[0].font.color.rgb = color_primary

    content = slide.placeholders[1]
    tf = content.text_frame
    
    p = tf.add_paragraph()
    p.text = "1. Proven Scalability: BOM < $25, ready for mass cluster deployment."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "2. Rural First: LoRaWAN connectivity handles 0-Wi-Fi environments."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "3. Export Ready: Directly aligns with India's Agri-Export National Policy."
    p.level = 1

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    import os
    target_path = os.path.join(os.getcwd(), "AgroSync_Pitch_Deck.pptx")
    create_presentation(target_path)
